from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.table_compare.api.handlers.excel_handler import ExcelHandler

from .run_helpers import make_saved_result


class ComparisonWorkbookTests(SimpleTestCase):
    def _load(self, result_data):
        output = ExcelHandler().build_comparison_workbook(result_data)
        self.assertEqual(output.tell(), 0)
        workbook = load_workbook(BytesIO(output.read()), data_only=True)
        output.close()
        self.addCleanup(workbook.close)
        return workbook

    def test_workbook_preserves_saved_rows_scope_issues_and_core_style(self):
        result_data = make_saved_result()
        workbook = self._load(result_data)

        self.assertEqual(
            workbook.sheetnames,
            ["比對結果", "比對範圍與解析提醒"],
        )
        result_sheet = workbook["比對結果"]
        self.assertEqual(result_sheet["A1"].value, "比對結果")
        self.assertIn(
            "A1:J1",
            {str(item) for item in result_sheet.merged_cells.ranges},
        )
        self.assertEqual(
            [cell.value for cell in result_sheet[2]],
            [
                "作物", "年度", "指標",
                "PDF單位", "PDF大項", "PDF數值",
                "Excel單位", "Excel大項", "Excel數值", "比對結果",
            ],
        )
        expected_rows = [
            [
                row["作物"], row["年度"], row["指標"],
                row["單位_pdf"], row["大項_pdf"], row["數值_pdf"],
                row["單位_excel"],
                row["大項_excel"],
                (
                    None
                    if row["數值_excel"] == ""
                    else row["數值_excel"]
                ),
                row["比對結果"],
            ]
            for row in result_data["rows"]
        ]
        self.assertEqual(
            [
                [cell.value for cell in row]
                for row in result_sheet.iter_rows(
                    min_row=3,
                    max_row=2 + result_data["row_count"],
                    min_col=1,
                    max_col=10,
                )
            ],
            expected_rows,
        )
        self.assertEqual(result_sheet["F3"].value, 0)
        self.assertEqual(result_sheet["I3"].value, "0")
        self.assertEqual(
            result_sheet["F4"].value,
            "123.12345678901234567890",
        )
        self.assertIsNone(result_sheet["I5"].value)
        self.assertEqual(result_sheet["F8"].value, "-")
        self.assertIsNone(result_sheet["I8"].value)
        self.assertEqual(
            [result_sheet.cell(row=index, column=10).value for index in range(3, 9)],
            [
                "相同", "數值不同", "Excel 缺少資料",
                "PDF 缺少資料", "單位不同", "缺值／標記，暫不比對",
            ],
        )
        self.assertEqual(result_sheet.freeze_panes, "D3")
        self.assertEqual(result_sheet["D2"].border.left.style, "thin")
        self.assertEqual(result_sheet["D3"].border.bottom.style, "thin")

        note_sheet = workbook["比對範圍與解析提醒"]
        self.assertEqual(note_sheet["A1"].value, "比對範圍與解析提醒")
        self.assertEqual(
            note_sheet["B3"].value,
            "香蕉、鳳梨、茶、芒果、番石榴、蓮霧",
        )
        self.assertEqual(note_sheet["B4"].value, "112")
        self.assertEqual(
            note_sheet["B5"].value,
            "產量、產值、種植面積、收穫面積",
        )
        self.assertEqual(note_sheet["B6"].value, 6)
        self.assertEqual(note_sheet["A8"].value, "來源解析提醒")
        self.assertEqual(note_sheet["B8"].value, "PDF 第 3 頁未解析")
        self.assertEqual(note_sheet["B9"].value, "茶未在 PDF 解析資料中找到")

    def test_workbook_distinguishes_no_issues_from_multiple_issues(self):
        workbook = self._load(make_saved_result(issues=[]))
        note_sheet = workbook["比對範圍與解析提醒"]
        self.assertEqual(note_sheet["A8"].value, "來源解析提醒")
        self.assertEqual(note_sheet["B8"].value, "目前沒有來源解析提醒。")

    def test_valid_empty_result_still_has_headers_and_saved_scope(self):
        result_data = {
            "row_count": 0,
            "rows": [],
            "scope": {
                "crops": ["香蕉"],
                "years": [112],
                "metrics": ["產量"],
            },
            "issues": [],
        }
        workbook = self._load(result_data)

        result_sheet = workbook["比對結果"]
        self.assertEqual(result_sheet.max_row, 2)
        self.assertEqual(result_sheet.max_column, 10)
        self.assertEqual(result_sheet["A2"].value, "作物")
        self.assertEqual(result_sheet["J2"].value, "比對結果")
        note_sheet = workbook["比對範圍與解析提醒"]
        self.assertEqual(note_sheet["B3"].value, "香蕉")
        self.assertEqual(note_sheet["B6"].value, 0)
        self.assertEqual(note_sheet["B8"].value, "目前沒有來源解析提醒。")
