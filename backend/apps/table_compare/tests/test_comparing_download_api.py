import json
import re
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.parse import unquote
from uuid import uuid4

from django.conf import settings
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from drf_spectacular.generators import SchemaGenerator
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITransactionTestCase

from apps.jobs.models import Job
from apps.table_compare.api.handlers.excel_handler import ExcelHandler
from apps.table_compare.api.services.comparing_service import ComparingService
from apps.table_compare.models import CompareInputFile

from .run_helpers import (
    DEFAULT_PREVIEW,
    create_compare_job,
    make_run_pdf_bytes,
    make_run_xlsx_bytes,
    make_saved_result,
)


class ComparingDownloadApiTests(APITransactionTestCase):
    XLSX_TYPE = (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )

    def setUp(self):
        media = TemporaryDirectory()
        self.media_root = Path(media.name)
        self.addCleanup(media.cleanup)
        media_settings = override_settings(MEDIA_ROOT=media.name)
        media_settings.enable()
        self.addCleanup(media_settings.disable)

    @staticmethod
    def download_url(public_id):
        return reverse(
            "table_compare:download",
            kwargs={"public_id": public_id},
        )

    @staticmethod
    def run_url(public_id):
        return reverse(
            "table_compare:comparing",
            kwargs={"public_id": public_id},
        )

    def get_download(self, public_id, *, query=""):
        return self.client.get(
            f"{self.download_url(public_id)}{query}",
            HTTP_ACCEPT=self.XLSX_TYPE,
        )

    def consume_workbook(self, response):
        body = b"".join(response.streaming_content)
        response.close()
        workbook = load_workbook(BytesIO(body), data_only=True)
        self.addCleanup(workbook.close)
        return body, workbook

    def assert_no_store(self, response):
        tokens = {
            token.strip().lower()
            for token in response.get("Cache-Control", "").split(",")
        }
        self.assertIn("no-store", tokens)

    @patch(
        "apps.table_compare.api.services.comparing_service.timezone.localdate",
        return_value=date(2026, 3, 5),
    )
    def test_succeeded_job_downloads_real_complete_workbook_and_headers(
        self,
        _localdate,
    ):
        result_data = make_saved_result()
        job, _compare_job = create_compare_job(
            status=Job.Status.SUCCEEDED,
            preview={
                **DEFAULT_PREVIEW,
                "issues": ["這是 preview，不可輸出"],
            },
            result_data=result_data,
        )

        response = self.get_download(job.public_id)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(settings.TIME_ZONE, "Asia/Taipei")
        self.assertEqual(response["Content-Type"], self.XLSX_TYPE)
        disposition = response["Content-Disposition"]
        self.assertIn("attachment", disposition.lower())
        match = re.search(r"filename\*=utf-8''([^;]+)", disposition, re.I)
        self.assertIsNotNone(match, disposition)
        self.assertEqual(
            unquote(match.group(1)),
            "115.03.05比對結果.xlsx",
        )
        self.assertNotIn("/app/", disposition)
        self.assertNotIn("\\", disposition)
        self.assert_no_store(response)

        body, workbook = self.consume_workbook(response)
        if "Content-Length" in response:
            self.assertEqual(int(response["Content-Length"]), len(body))
        result_sheet = workbook["比對結果"]
        self.assertEqual(result_sheet.max_row, result_data["row_count"] + 2)
        self.assertEqual(result_sheet["A3"].value, "香蕉")
        self.assertEqual(result_sheet["A8"].value, "蓮霧")
        note_sheet = workbook["比對範圍與解析提醒"]
        self.assertEqual(note_sheet["B8"].value, result_data["issues"][0])
        self.assertNotIn(
            "這是 preview，不可輸出",
            [
                cell.value
                for row in note_sheet.iter_rows()
                for cell in row
            ],
        )

    def test_not_found_and_other_kind_are_404_without_mutation(self):
        other_job = Job.objects.create(
            kind=Job.Kind.VOLUME_PRICE_MERGE,
            status=Job.Status.SUCCEEDED,
            expires_at=timezone.now() + timedelta(hours=1),
        )
        original_updated_at = other_job.updated_at

        for public_id in (uuid4(), other_job.public_id):
            with self.subTest(public_id=public_id):
                response = self.get_download(public_id)
                self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
                self.assertEqual(response.json()["code"], "job_not_found")
                self.assertEqual(
                    set(response.json()),
                    {"code", "message", "field_errors"},
                )
                self.assert_no_store(response)

        other_job.refresh_from_db()
        self.assertEqual(other_job.status, Job.Status.SUCCEEDED)
        self.assertEqual(other_job.updated_at, original_updated_at)

    def test_non_succeeded_jobs_are_409_without_export_or_mutation(self):
        result_data = make_saved_result()
        with patch.object(
            ExcelHandler,
            "build_comparison_workbook",
        ) as exporter:
            for job_status in (
                Job.Status.PENDING,
                Job.Status.RUNNING,
                Job.Status.FAILED,
            ):
                with self.subTest(job_status=job_status):
                    job, compare_job = create_compare_job(
                        status=job_status,
                        result_data=result_data,
                    )
                    original_updated_at = job.updated_at
                    response = self.get_download(job.public_id)

                    self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
                    self.assertEqual(response.json()["code"], "job_not_runnable")
                    self.assertIn("下載", response.json()["message"])
                    self.assert_no_store(response)
                    job.refresh_from_db()
                    compare_job.refresh_from_db()
                    self.assertEqual(job.status, job_status)
                    self.assertEqual(job.updated_at, original_updated_at)
                    self.assertEqual(compare_job.result_data, result_data)
        exporter.assert_not_called()

    def test_real_run_persists_then_downloads_without_mocking_core_path(self):
        job, _compare_job = create_compare_job(
            with_files=True,
            excel_content=make_run_xlsx_bytes(value=10),
            pdf_content=make_run_pdf_bytes(value=10),
        )

        run_response = self.client.post(
            self.run_url(job.public_id),
            DEFAULT_PREVIEW["defaults"],
            format="json",
        )
        self.assertEqual(
            run_response.status_code,
            status.HTTP_200_OK,
            run_response.json(),
        )

        download_response = self.get_download(job.public_id)

        self.assertEqual(download_response.status_code, status.HTTP_200_OK)
        _body, workbook = self.consume_workbook(download_response)
        self.assertEqual(workbook["比對結果"]["A3"].value, "香蕉")
        self.assertEqual(workbook["比對結果"]["J3"].value, "相同")
        job.refresh_from_db()
        self.assertEqual(job.status, Job.Status.SUCCEEDED)

    def test_download_schema_declares_binary_get_and_error_contracts(self):
        schema = SchemaGenerator().get_schema(request=None, public=True)

        def resolve(value):
            while "$ref" in value:
                value = schema["components"]["schemas"][
                    value["$ref"].split("/")[-1]
                ]
            return value

        operation = schema["paths"][
            "/api/v1/table-compare/jobs/{public_id}/download/"
        ]["get"]
        self.assertEqual(operation["summary"], "下載表格比對結果")
        self.assertNotIn("requestBody", operation)
        parameters = operation.get("parameters", [])
        self.assertEqual(
            [(item["name"], item["in"]) for item in parameters],
            [("public_id", "path")],
        )
        self.assertEqual(parameters[0]["schema"]["format"], "uuid")
        self.assertEqual(
            set(operation["responses"]),
            {"200", "404", "409", "410", "500"},
        )
        binary = operation["responses"]["200"]["content"][
            self.XLSX_TYPE
        ]["schema"]
        self.assertEqual(binary["type"], "string")
        self.assertEqual(binary["format"], "binary")
        for response_code in ("404", "409", "410", "500"):
            error = resolve(
                operation["responses"][response_code]["content"][
                    "application/json"
                ]["schema"]
            )
            self.assertEqual(
                set(error["properties"]),
                {"code", "message", "field_errors"},
            )

    def test_missing_or_malformed_saved_result_is_controlled_409(self):
        missing_relation = Job.objects.create(
            kind=Job.Kind.TABLE_COMPARE,
            status=Job.Status.SUCCEEDED,
            expires_at=timezone.now() + timedelta(hours=1),
        )
        malformed_results = (
            None,
            {},
            {
                "row_count": 0,
                "rows": [],
                "scope": {"crops": [], "years": [], "metrics": []},
            },
            {
                "row_count": 1,
                "rows": [{"作物": "香蕉"}],
                "scope": {
                    "crops": ["香蕉"],
                    "years": [112],
                    "metrics": ["產量"],
                },
                "issues": [],
            },
        )
        public_ids = [missing_relation.public_id]
        for result_data in malformed_results:
            job, _compare_job = create_compare_job(
                status=Job.Status.SUCCEEDED,
                result_data=result_data,
            )
            public_ids.append(job.public_id)

        with patch.object(
            ExcelHandler,
            "build_comparison_workbook",
        ) as exporter:
            for public_id in public_ids:
                with self.subTest(public_id=public_id):
                    response = self.get_download(public_id)
                    self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
                    self.assertEqual(response.json()["code"], "job_not_runnable")
        exporter.assert_not_called()

    def test_valid_empty_saved_result_downloads_headers(self):
        result_data = {
            "row_count": 0,
            "rows": [],
            "scope": {
                "crops": ["香蕉"],
                "years": [112],
                "metrics": ["產量"],
            },
            "issues": [],
        }
        job, _compare_job = create_compare_job(
            status=Job.Status.SUCCEEDED,
            result_data=result_data,
        )

        response = self.get_download(job.public_id)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        _body, workbook = self.consume_workbook(response)
        self.assertEqual(workbook["比對結果"].max_row, 2)
        self.assertEqual(workbook["比對結果"]["J2"].value, "比對結果")

    def test_export_failure_is_safe_500_and_does_not_mutate_job(self):
        result_data = make_saved_result()
        job, compare_job = create_compare_job(
            status=Job.Status.SUCCEEDED,
            result_data=result_data,
        )
        original_updated_at = job.updated_at
        secret = "C:/private/internal/export-secret"

        with patch.object(
            ExcelHandler,
            "build_comparison_workbook",
            side_effect=RuntimeError(secret),
        ):
            response = self.get_download(job.public_id)

        self.assertEqual(
            response.status_code,
            status.HTTP_500_INTERNAL_SERVER_ERROR,
        )
        self.assertEqual(response.json()["code"], "internal_server_error")
        serialized = json.dumps(response.json(), ensure_ascii=False)
        self.assertNotIn(secret, serialized)
        self.assertNotIn("Traceback", serialized)
        self.assert_no_store(response)
        job.refresh_from_db()
        compare_job.refresh_from_db()
        self.assertEqual(job.status, Job.Status.SUCCEEDED)
        self.assertEqual(job.updated_at, original_updated_at)
        self.assertEqual(compare_job.result_data, result_data)

    def test_repeated_get_and_product_query_keep_full_result_and_no_side_effects(self):
        result_data = make_saved_result()
        job, compare_job = create_compare_job(
            status=Job.Status.SUCCEEDED,
            result_data=result_data,
            with_files=True,
        )
        original_updated_at = job.updated_at
        original_files = list(
            CompareInputFile.objects.filter(compare_job=compare_job)
            .order_by("pk")
            .values_list("file_type", "file", "original_name", "size_bytes")
        )
        disk_files_before = sorted(
            path.relative_to(self.media_root)
            for path in self.media_root.rglob("*")
            if path.is_file()
        )

        with (
            patch.object(ComparingService, "run") as run,
            patch(
                "apps.table_compare.api.services.comparing_service."
                "ExcelTableAdapter.load"
            ) as excel_load,
            patch(
                "apps.table_compare.api.services.comparing_service."
                "PDFTableAdapter.load"
            ) as pdf_load,
        ):
            first = self.get_download(job.public_id)
            _first_body, first_workbook = self.consume_workbook(first)
            second = self.get_download(
                job.public_id,
                query="?product=香蕉",
            )
            _second_body, second_workbook = self.consume_workbook(second)

        run.assert_not_called()
        excel_load.assert_not_called()
        pdf_load.assert_not_called()
        first_rows = [
            [cell.value for cell in row]
            for row in first_workbook["比對結果"].iter_rows(
                min_row=3,
                max_col=10,
            )
        ]
        second_rows = [
            [cell.value for cell in row]
            for row in second_workbook["比對結果"].iter_rows(
                min_row=3,
                max_col=10,
            )
        ]
        self.assertEqual(first_rows, second_rows)
        self.assertEqual(len(second_rows), result_data["row_count"])
        job.refresh_from_db()
        compare_job.refresh_from_db()
        self.assertEqual(job.status, Job.Status.SUCCEEDED)
        self.assertEqual(job.updated_at, original_updated_at)
        self.assertEqual(job.error_code, "")
        self.assertEqual(job.error_message, "")
        self.assertEqual(compare_job.result_data, result_data)
        self.assertEqual(
            list(
                CompareInputFile.objects.filter(compare_job=compare_job)
                .order_by("pk")
                .values_list("file_type", "file", "original_name", "size_bytes")
            ),
            original_files,
        )
        disk_files_after = sorted(
            path.relative_to(self.media_root)
            for path in self.media_root.rglob("*")
            if path.is_file()
        )
        self.assertEqual(disk_files_after, disk_files_before)

    def test_expired_jobs_are_410_and_time_expiry_is_persisted(self):
        cases = (
            create_compare_job(
                status=Job.Status.SUCCEEDED,
                expired=True,
                result_data=make_saved_result(),
            ),
            create_compare_job(
                status=Job.Status.EXPIRED,
                result_data=make_saved_result(),
            ),
        )
        with patch.object(
            ExcelHandler,
            "build_comparison_workbook",
        ) as exporter:
            for job, compare_job in cases:
                with self.subTest(job_status=job.status):
                    original_result = compare_job.result_data
                    response = self.get_download(job.public_id)
                    self.assertEqual(response.status_code, status.HTTP_410_GONE)
                    self.assertEqual(response.json()["code"], "job_expired")
                    self.assert_no_store(response)
                    job.refresh_from_db()
                    compare_job.refresh_from_db()
                    self.assertEqual(job.status, Job.Status.EXPIRED)
                    self.assertEqual(compare_job.result_data, original_result)
        exporter.assert_not_called()
