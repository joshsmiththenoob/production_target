"""
Duty on every jobs in Excel, which has composition of workbook reader 
and some other excel application objects.
"""
import io
from datetime import datetime
from typing import Any
from collections import OrderedDict
from openpyxl import Workbook


from .re_handler import REHandler
from .excel_styler import ExcelStyler


class ExcelHandler:
    def __init__(self):
        self.__RE_handler = REHandler()
        self.__excel_styler = ExcelStyler()


    def _detect_header_matrix_structure(self, rows) -> dict:
        scan_limit = min(len(rows), 8)
        # Find index of year header
        year_row_index = [
            index for index, row in enumerate(rows[:scan_limit])
            if any(self.__RE_handler.is_year(cell) for cell in row)
        ]
        if not year_row_index:
         raise ValueError("Can't find any year information in excel file. Please check header structure of file.")
        

        # Find index of crop/product's header 
        start = max(year_row_index) + 1
        candidates: list[tuple[int, int]] = []

        
        for row_index in range(start, scan_limit):
            row = rows[row_index]
            score = 0
            for value in row[1:]:
                if value and not self.__RE_handler.is_year(value) and value not in {"合計", "總計", "品項", "作物"}:
                    score += 1
            if score:
                candidates.append((score, row_index))
        if not candidates:
            raise ValueError("找不到作物列，請確認年度列下方是否有品項名稱。")
        crop_row_index = max(candidates, key=lambda item: (item[0], -item[1]))[1]


        # Find index of starting of dataset
        data_start_index = None
        for row_index in range(crop_row_index + 1, len(rows)):
            row = rows[row_index]
            if row and not self.__RE_handler.is_blank(row[0]) and any(not self.__RE_handler.is_blank(cell) for cell in row[1:]):
                data_start_index = row_index
                break
        if data_start_index is None:
            raise ValueError("找不到統計指標資料列。")

        return {
            "year_rows": year_row_index,
            "crop_row_index": crop_row_index,
            "data_start_index": data_start_index,
        }

    
    def normalize_input(self, rows: list[list[Any]], major_category: str, source_type: str, file_name: str) -> dict[str, Any]:
        header = self._detect_header_matrix_structure(rows)
        max_columns = max((len(row) for row in rows), default=0)
        columns: list[dict[str, Any]] = []

        # Get current_year
        current_year = ""
        for col_index in range(1, max_columns):
            for year_row_index in header["year_rows"]:
                value = rows[year_row_index][col_index] if col_index < len(rows[year_row_index]) else None
                if self.__RE_handler.is_year(value):
                    current_year = value


            crop = rows[header["crop_row_index"]][col_index] if col_index < len(rows[header["crop_row_index"]]) else ""
            if current_year and crop:
                columns.append({
                    "majorCategory": major_category,
                    "year": current_year,
                    "crop": crop,
                    "sourceCol": col_index,
                    "key": self.column_key(major_category, current_year, crop),
                })
        if not columns:
            raise ValueError(f"「{file_name}」表頭無法對齊，找不到年度與作物欄位。")

        output_rows: list[dict[str, Any]] = []
        for row in rows[header["data_start_index"]:]:
            metric = row[0] if row else ""
            if not metric:
                continue
            values: OrderedDict[str, list[Any]] = OrderedDict()
            for column in columns:
                value = row[column["sourceCol"]] if column["sourceCol"] < len(row) else None
                values.setdefault(column["key"], []).append(value)
            output_rows.append({"metric": metric, "valuesByKey": dict(values)})
        if not output_rows:
            raise ValueError(f"「{file_name}」找不到有效的統計指標資料列。")
        return {
            "majorCategory": major_category,
            "sourceType": source_type,
            "fileName": file_name,
            "columns": columns,
            "rows": output_rows,
        }


    def column_key(self, major_category: str, year: str, crop: str) -> str:
        return f"{major_category}||{year}||{crop}"


    def build_column_order(self, all_sources: list[dict[str, Any]]) -> list[dict[str, str]]:
        seen: set[str] = set()
        major_order: OrderedDict[str, int] = OrderedDict()
        crop_order: dict[tuple[str, str], int] = {}
        result = []
        for source in all_sources:
            major = source["majorCategory"]
            major_order.setdefault(major, len(major_order))
            for column in source["columns"]:
                crop_order.setdefault((major, column["crop"]), len(crop_order))
                if column["key"] not in seen:
                    seen.add(column["key"])
                    result.append({key: column[key] for key in ("key", "majorCategory", "year", "crop")})
        
        
        return sorted(result, key=lambda item: (
            major_order[item["majorCategory"]],
            self.__RE_handler.year_number(item["year"]),
            crop_order[(item["majorCategory"], item["crop"])],
        ))


    def build_merged_rows(self, all_sources: list[dict[str, Any]], column_order: list[dict[str, str]]) -> list[dict[str, Any]]:
        merged_rows = []
        for source in all_sources:
            for source_row in source["rows"]:
                merged_rows.append({
                    "metric": source_row["metric"],
                    "sourceType": source["sourceType"],
                    "majorCategory": source["majorCategory"],
                    "values": [self._collapse_duplicate_values(source_row["valuesByKey"].get(column["key"], [])) for column in column_order],
                })
        return merged_rows


    def _collapse_duplicate_values(self, values: list[Any]) -> Any:
        if not values:
            return None
        if len(values) == 1:
            return values[0]
        return " / ".join( "" if self.__RE_handler.is_blank(value) else str(value) for value in values)


    def build_result_data(self, column_order: list[dict[str, Any]], merged_rows: list[dict[str, Any]]) -> dict[str, Any]:
        """
        Build results from camelCase to snake_case
        """
        
        columns = [{
            "key": column["key"],
            "major_category": column["majorCategory"],
            "year": column["year"],
            "crop": column["crop"]
        } for column in column_order]


        rows = [
            {
            "metric": row["metric"],
            "property_type": row["sourceType"],
            "major_category": row["majorCategory"],
            "values": row["values"],
        }for row in merged_rows
        ]

        # Extract deduplicated product(crop) name from columns
        # A sorted set matches the prototype's stable crop options for Step 4.
        available_crops = sorted({column["crop"] for column in columns})

        return {
            "schema_version": 1,
            "columns": columns,
            "rows": rows,
            "available_crops": available_crops,
            "summary": {
                "column_count": len(columns),
                "row_count": len(rows),
            }
        }


    # def query(job_id: str):
    #     """
    #     Query the specific product in result of specialized job
    #     """
    #     pass


    def filter_by_product(self, result_data: dict, product: str= '') -> dict:
        """
        Find specific product's merged result
        """
        if not product:
            raise ValueError("Doesn't choose product yet. Please choose 1 product.")
        
        
        matched = [(index, column) for index, column in enumerate(result_data["columns"]) if column["crop"] == product]
        if not matched:
           raise ValueError(f"Invalid product name 「{product}」 in result of job. -> Can't find this product")

        
        result_rows = []
        for row in result_data["rows"]:
            values = [row["values"][index] for index, _column in matched]
            if any(not self.__RE_handler.is_blank(value) for value in values):
                result_rows.append({"metric": row["metric"], "values": values})


        if not result_rows:
            raise ValueError(f"Invalid「{product}」data in this result of job")

        
        return {
                "product": product, 
                "columns": [column for _index, column in matched], 
                "rows": result_rows
                }


    def _merge_header_cells(self, sheet, row_index: int, values: list[str]) -> None:
        start = 0
        while start < len(values):
            end = start
            while end + 1 < len(values) and values[end + 1] == values[start]:
                end += 1
            if end > start:
                sheet.merge_cells(start_row=row_index, start_column=start + 2, end_row=row_index, end_column=end + 2)
            start = end + 1


    def build_workbook(self, all_sources: list[dict[str, Any]], column_order: list[dict[str, str]], merged_rows: list[dict[str, Any]], validation: dict[str, Any]) -> Workbook:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "合併結果"
        total_columns = len(column_order) + 1
        rows = [
            ["合併結果"] + [None] * (total_columns - 1),
            ["統計指標"] + [column["year"] for column in column_order],
            [None] + [column["majorCategory"] for column in column_order],
            [None] + [column["crop"] for column in column_order],
        ]
        rows.extend([[row["metric"]] + row["values"] for row in merged_rows])
        for row in rows:
            sheet.append(row)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
        sheet.merge_cells(start_row=2, start_column=1, end_row=4, end_column=1)
        self._merge_header_cells(sheet, 2, [column["year"] for column in column_order])
        self._merge_header_cells(sheet, 3, [column["majorCategory"] for column in column_order])
        self.__excel_styler.apply_table_style(sheet, len(rows), total_columns)

        source_sheet = workbook.create_sheet("來源說明")
        source_rows = [
            ["項目", "內容"],
            ["合併主鍵", "左側以統計指標保留並對齊；欄位唯一鍵為大項 + 年度 + 作物。"],
            ["來源檔案數", len(all_sources)],
            ["完整大項", "、".join(pair["majorCategory"] for pair in validation["complete"])],
            ["保留方式", "重複統計指標與重複值全部保留，不去重；同一列多值以 / 串接。"],
        ]
        for row in source_rows:
            source_sheet.append(row)
        self.__excel_styler.style_note_sheet(source_sheet, len(source_rows))

        pair_sheet = workbook.create_sheet("配對結果")
        pair_sheet.append(["大項", "產量及產值檔案", "種植及收穫面積檔案", "結果"])
        for pair in validation["complete"] + validation["incomplete"]:
            pair_sheet.append([
                pair["majorCategory"],
                "\n".join(item["fileName"] for item in pair["production"]) or "缺少",
                "\n".join(item["fileName"] for item in pair["area"]) or "缺少",
                "完整" if pair["production"] and pair["area"] else "不完整",
            ])
        pair_sheet.column_dimensions["A"].width = 20
        pair_sheet.column_dimensions["B"].width = 48
        pair_sheet.column_dimensions["C"].width = 48
        pair_sheet.column_dimensions["D"].width = 14
        self.__excel_styler.style_note_sheet(pair_sheet, pair_sheet.max_row, 4)
        return workbook


    def build_query_workbook(self, job: dict[str, Any], query: dict[str, Any]) -> Workbook:
        columns = query["columns"]
        rows = query["rows"]
        total_columns = len(columns) + 1
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "查詢結果"
        table = [
            [f"合併結果－作物查詢：{query['crop']}"] + [None] * (total_columns - 1),
            ["統計指標"] + [column["majorCategory"] for column in columns],
            [None] + [column["year"] for column in columns],
            [None] + [column["crop"] for column in columns],
        ]
        table.extend([[row["metric"]] + row["values"] for row in rows])
        for row in table:
            sheet.append(row)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
        sheet.merge_cells(start_row=2, start_column=1, end_row=4, end_column=1)
        self._merge_header_cells(sheet, 2, [column["majorCategory"] for column in columns])
        self._merge_header_cells(sheet, 3, [column["year"] for column in columns])
        self.__excel_styler.apply_table_style(sheet, len(table), total_columns)

        condition = workbook.create_sheet("查詢條件")
        conditions = [
            ["項目", "內容"],
            ["查詢作物", query["crop"]],
            ["下載時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["來源", "來自目前已合併資料"],
            ["欄位篩選", "僅保留所選作物欄位"],
            ["列篩選", "僅保留至少一個非空值的統計指標列"],
        ]
        for row in conditions:
            condition.append(row)
        self.__excel_styler.style_note_sheet(condition, len(conditions))
        return workbook


    def _workbook_bytes(self, workbook: Workbook) -> io.ByteIO:
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
