"""
The common Template of Excel Styler for reuse
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter



class ExcelStyler:
    def __init__(self):
        pass

    def apply_table_style(self, sheet, total_rows: int, total_columns: int) -> None:
        border_side = Side(style="thin", color="D0D7DE")
        border = Border(top=border_side, bottom=border_side, left=border_side, right=border_side)
        for row in sheet.iter_rows(min_row=1, max_row=total_rows, min_col=1, max_col=total_columns):
            for cell in row:
                cell.font = Font(name="Microsoft JhengHei")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF", size=14)
        for row in sheet.iter_rows(min_row=2, max_row=min(4, total_rows), min_col=1, max_col=total_columns):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(name="Microsoft JhengHei", bold=True, color="17365D")
        for row_index in range(5, total_rows + 1):
            cell = sheet.cell(row=row_index, column=1)
            cell.fill = PatternFill("solid", fgColor="F6F8FA")
            cell.font = Font(name="Microsoft JhengHei", bold=True)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.column_dimensions["A"].width = 22
        for index in range(2, total_columns + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 15
        for index in range(1, min(total_rows, 4) + 1):
            sheet.row_dimensions[index].height = 24
        sheet.freeze_panes = "B5"
        sheet.sheet_view.showGridLines = False



    def style_note_sheet(self, sheet, total_rows: int, total_columns: int = 2) -> None:
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 100
        for row in sheet.iter_rows(min_row=1, max_row=total_rows, min_col=1, max_col=total_columns):
            for cell in row:
                cell.font = Font(name="Microsoft JhengHei")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF")