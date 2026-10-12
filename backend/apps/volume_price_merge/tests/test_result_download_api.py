from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.jobs.models import Job
from apps.volume_price_merge.models import VolumePriceMergeJob


class ResultAndDownloadApiTests(APITestCase):
    def _create_job(
        self,
        *,
        status_value: str = Job.Status.SUCCEEDED,
        expired: bool = False,
        result_data: dict | None = None,
    ) -> Job:
        job = Job.objects.create(
            kind=Job.Kind.VOLUME_PRICE_MERGE,
            status=status_value,
            expires_at=timezone.now() + timedelta(hours=-1 if expired else 1),
        )
        VolumePriceMergeJob.objects.create(
            job=job,
            pairing_preview={"pairs": []},
            result_data=result_data or self._result_data(),
        )
        return job

    @staticmethod
    def _result_data() -> dict:
        return {
            "schema_version": 1,
            "columns": [
                {"key": "果品||110年||香蕉", "major_category": "果品", "year": "110年", "crop": "香蕉"},
                {"key": "果品||111年||香蕉", "major_category": "果品", "year": "111年", "crop": "香蕉"},
                {"key": "果品||110年||鳳梨", "major_category": "果品", "year": "110年", "crop": "鳳梨"},
            ],
            "rows": [
                {"metric": "果品產量", "property_type": "production", "major_category": "果品", "values": [0, "10 / 20", 99]},
                {"metric": "果品產量", "property_type": "area", "major_category": "果品", "values": [None, 5, 88]},
            ],
            "available_crops": ["香蕉", "鳳梨"],
            "summary": {"column_count": 3, "row_count": 2},
        }

    @staticmethod
    def _query_url(public_id) -> str:
        return reverse("volume_price_merge:query", kwargs={"public_id": public_id})

    @staticmethod
    def _download_url(public_id) -> str:
        return reverse("volume_price_merge:download", kwargs={"public_id": public_id})

    def assert_no_store(self, response) -> None:
        self.assertIn("no-store", response.headers["Cache-Control"])

    def test_query_preserves_order_duplicate_metrics_and_scalar_values(self) -> None:
        job = self._create_job()

        response = self.client.get(self._query_url(job.public_id), {"product": "香蕉"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assert_no_store(response)
        payload = response.json()["data"]
        self.assertEqual(payload["public_id"], str(job.public_id))
        self.assertEqual(payload["query_result"]["product"], "香蕉")
        self.assertEqual(
            [column["year"] for column in payload["query_result"]["columns"]],
            ["110年", "111年"],
        )
        self.assertEqual(
            payload["query_result"]["rows"],
            [
                {"metric": "果品產量", "values": [0, "10 / 20"]},
                {"metric": "果品產量", "values": [None, 5]},
            ],
        )

    def test_query_allows_a_valid_empty_result(self) -> None:
        data = self._result_data()
        data["rows"] = [
            {"metric": "果品產量", "property_type": "production", "major_category": "果品", "values": [None, None, 99]},
        ]
        job = self._create_job(result_data=data)

        response = self.client.get(self._query_url(job.public_id), {"product": "香蕉"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["data"]["query_result"]["rows"], [])

    def test_result_and_download_input_errors_are_400_and_no_store(self) -> None:
        job = self._create_job()

        missing_product = self.client.get(self._query_url(job.public_id))
        unknown_product = self.client.get(self._download_url(job.public_id), {"product": "不存在"})

        for response in (missing_product, unknown_product):
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertEqual(response.json()["code"], "invalid_product")
            self.assert_no_store(response)

    def test_result_and_download_job_errors_are_explicit_and_no_store(self) -> None:
        cases = [
            (uuid4(), status.HTTP_404_NOT_FOUND, "job_not_found"),
            (self._create_job(status_value=Job.Status.PENDING).public_id, status.HTTP_409_CONFLICT, None),
            (self._create_job(expired=True).public_id, status.HTTP_410_GONE, "job_expired"),
        ]

        for public_id, expected_status, expected_code in cases:
            for route_name in ("query", "download"):
                with self.subTest(public_id=public_id, route=route_name):
                    url = reverse(f"volume_price_merge:{route_name}", kwargs={"public_id": public_id})
                    response = self.client.get(url, {"product": "香蕉"})
                    self.assertEqual(response.status_code, expected_status)
                    if expected_code is not None:
                        self.assertEqual(response.json()["code"], expected_code)
                    self.assert_no_store(response)

    def test_full_and_filtered_downloads_are_valid_xlsx_files(self) -> None:
        job = self._create_job()

        full_response = self.client.get(self._download_url(job.public_id))
        filtered_response = self.client.get(self._download_url(job.public_id), {"product": "香蕉"})

        for response in (full_response, filtered_response):
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertEqual(
                response.headers["Content-Type"],
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.assert_no_store(response)

        full_workbook = load_workbook(BytesIO(b"".join(full_response.streaming_content)), data_only=True)
        self.assertEqual(full_workbook["合併結果"].cell(5, 2).value, 0)
        self.assertEqual(full_workbook["合併結果"].cell(5, 3).value, "10 / 20")

        filtered_workbook = load_workbook(BytesIO(b"".join(filtered_response.streaming_content)), data_only=True)
        sheet = filtered_workbook["查詢結果"]
        self.assertEqual(sheet.cell(1, 1).value, "合併結果－作物查詢：香蕉")
        self.assertEqual(sheet.cell(5, 1).value, "果品產量")
        self.assertEqual(sheet.cell(6, 1).value, "果品產量")
        self.assertEqual(sheet.cell(5, 2).value, 0)
        self.assertEqual(sheet.cell(5, 3).value, "10 / 20")
